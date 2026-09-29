"""Testler için ortak ayarlar ve örnek dosya üreticileri.

Örnek dosyalar dönüştürücünün kendi yazıcılarıyla DEĞİL, bağımsız
kütüphanelerle (python-docx, reportlab, odfpy, Pillow, ffmpeg) üretilir;
böylece okuyucular gerçekçi girdilerle sınanır.
"""

import html
import io
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

KOK = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KOK))

from converters import belge  # noqa: E402
from converters.belge_bloklar import metin_coz  # noqa: E402

WORD_VAR = belge._word_kullanilabilir()
word_gerekli = pytest.mark.skipif(not WORD_VAR, reason="Microsoft Word kurulu değil")

# Türkçe karakter denetimi için her örnekte bulunan işaret metni
ISARET = "Şişli Çağrı Öğün"


def png_bayt(boyut=(60, 40), renk="red") -> bytes:
    from PIL import Image

    tampon = io.BytesIO()
    Image.new("RGB", boyut, renk).save(tampon, "PNG")
    return tampon.getvalue()


def _rtf_kacir(metin: str) -> str:
    return "".join(c if ord(c) < 128 else f"\\u{ord(c)}?" for c in metin)


# ============================================================ belge örnekleri

def _docx_uret(yol: Path):
    import docx
    from docx.oxml import OxmlElement
    from docx.shared import Cm

    d = docx.Document()
    d.add_heading("Başlık Bir", level=1)
    p = d.add_paragraph("Normal ")
    p.add_run("kalın").bold = True
    p.add_run(" ve ")
    p.add_run("eğik").italic = True
    d.add_paragraph(ISARET)
    d.add_paragraph("madde bir", style="List Bullet")
    d.add_paragraph("madde iki", style="List Bullet")
    d.add_picture(io.BytesIO(png_bayt()), width=Cm(3))
    tablo = d.add_table(rows=2, cols=2)
    for i, satir in enumerate(tablo.rows):
        for j, hucre in enumerate(satir.cells):
            hucre.text = f"h{i}{j}"
    # köprü içindeki metin (eskiden kayboluyordu)
    p = d.add_paragraph("Bağlantı: ")
    kopru = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "KÖPRÜMETNİ"
    run.append(t)
    kopru.append(run)
    p._p.append(kopru)
    d.save(yol)


def _pdf_uret(yol: Path):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    if "TestArial" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("TestArial", r"C:\Windows\Fonts\arial.ttf"))
    c = canvas.Canvas(str(yol), pagesize=A4)
    c.setFont("TestArial", 18)
    c.drawString(72, 760, "Başlık Bir")
    c.setFont("TestArial", 12)
    c.drawString(72, 730, ISARET)
    c.drawString(72, 710, "İkinci satır metni")
    c.save()


def _odt_uret(yol: Path):
    from odf import text as odf_text
    from odf.opendocument import OpenDocumentText
    from odf.style import Style, TextProperties
    from odf.table import Table, TableCell, TableColumn, TableRow

    d = OpenDocumentText()
    kalin = Style(name="K", family="text")
    kalin.addElement(TextProperties(fontweight="bold"))
    d.automaticstyles.addElement(kalin)
    d.text.addElement(odf_text.H(outlinelevel=1, text="Başlık Bir"))
    p = odf_text.P(text="Normal ")
    p.addElement(odf_text.Span(stylename=kalin, text="kalın"))
    d.text.addElement(p)
    bolum = odf_text.Section(name="Bolum1")  # bölüm içeriği eskiden atlanıyordu
    bolum.addElement(odf_text.P(text=ISARET))
    d.text.addElement(bolum)
    liste = odf_text.List()
    for madde in ("madde bir", "madde iki"):
        oge = odf_text.ListItem()
        oge.addElement(odf_text.P(text=madde))
        liste.addElement(oge)
    d.text.addElement(liste)
    from odf import draw

    href = d.addPictureFromString(png_bayt(), "image/png")
    cerceve = draw.Frame(width="3cm", height="2cm", anchortype="as-char")
    cerceve.addElement(draw.Image(href=href))
    p = odf_text.P()
    p.addElement(cerceve)
    d.text.addElement(p)
    tablo = Table(name="T1")
    tablo.addElement(TableColumn(numbercolumnsrepeated=2))
    for i in range(2):
        satir = TableRow()
        for j in range(2):
            hucre = TableCell()
            hucre.addElement(odf_text.P(text=f"h{i}{j}"))
            satir.addElement(hucre)
        tablo.addElement(satir)
    d.text.addElement(tablo)
    d.save(str(yol))


_HTML = (
    '<!DOCTYPE html><html><head><meta charset="utf-8"><title>t</title></head><body>'
    "<h1>Başlık Bir</h1><p>Normal <b>kalın</b> ve <i>eğik</i></p>"
    f"<p>{ISARET}</p><ul><li>madde bir</li><li>madde iki</li></ul>"
    '<p><img src="resim.png" alt="r"></p>'
    "<table><tr><td>h00</td><td>h01</td></tr></table></body></html>"
)
_MD = (
    "# Başlık Bir\n\nNormal **kalın** ve *eğik*\n\n"
    f"{ISARET}\n\n- madde bir\n- madde iki\n\n![resim](resim.png)\n\n"
    "| a | b |\n|---|---|\n| h00 | h01 |\n"
)


@pytest.fixture(scope="session")
def belge_ornekleri(tmp_path_factory) -> dict[str, Path]:
    kl = tmp_path_factory.mktemp("belge_ornek")
    (kl / "resim.png").write_bytes(png_bayt())
    ornek = {f: kl / f"ornek.{f}" for f in ("pdf", "docx", "txt", "html", "md", "rtf", "odt")}
    _docx_uret(ornek["docx"])
    _pdf_uret(ornek["pdf"])
    _odt_uret(ornek["odt"])
    ornek["html"].write_text(_HTML, encoding="utf-8")
    ornek["md"].write_text(_MD, encoding="utf-8")
    ornek["txt"].write_text(f"Başlık Bir\n{ISARET}\nüçüncü satır\n", encoding="utf-8")
    ornek["rtf"].write_text(
        "{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}}\n"
        f"\\pard\\b {_rtf_kacir('Başlık Bir')}\\b0\\par\n"
        f"\\pard {_rtf_kacir(ISARET)}\\par\n}}", encoding="ascii")
    if WORD_VAR:
        ornek["doc"] = kl / "ornek.doc"
        _word_kaydet(ornek["docx"], ornek["doc"], 0)  # wdFormatDocument97
    return ornek


def _word_kaydet(kaynak: Path, hedef: Path, bicim: int):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    try:
        word.Visible = False
        word.DisplayAlerts = 0
        d = word.Documents.Open(str(kaynak), False, True, False)
        d.SaveAs2(str(hedef), FileFormat=bicim)
        d.Close(0)
    finally:
        word.Quit()
        pythoncom.CoUninitialize()


# ============================================================ çıktı okuyucu

def metni_oku(yol: Path) -> str:
    """Bir çıktı dosyasındaki düz metni (biçimden bağımsız) döndürür."""
    uzanti = yol.suffix.lower().lstrip(".")
    if uzanti in ("txt", "md"):
        metin = yol.read_text(encoding="utf-8")
        metin = metin.replace("\\", "")  # MD kaçışları
    elif uzanti == "html":
        metin = metin_coz(yol.read_bytes(), html_mi=True)
        metin = re.sub(r"(?is)<(script|style|head)\b.*?</\1>", " ", metin)
        metin = html.unescape(re.sub(r"<[^>]+>", " ", metin))
    elif uzanti == "docx":
        with zipfile.ZipFile(yol) as z:
            xml = z.read("word/document.xml").decode("utf-8")
        metin = html.unescape("".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", xml)))
    elif uzanti == "odt":
        with zipfile.ZipFile(yol) as z:
            xml = z.read("content.xml").decode("utf-8")
        metin = html.unescape(re.sub(r"<[^>]+>", "", xml.replace("<text:s/>", " ")))
    elif uzanti == "pdf":
        from pypdf import PdfReader

        metin = "\n".join(s.extract_text() or "" for s in PdfReader(str(yol)).pages)
    elif uzanti == "rtf":
        from striprtf.striprtf import rtf_to_text

        metin = rtf_to_text(metin_coz(yol.read_bytes()), errors="replace")
    else:
        raise ValueError(uzanti)
    return " ".join(metin.split())


def resim_var_mi(yol: Path) -> bool:
    uzanti = yol.suffix.lower().lstrip(".")
    if uzanti == "docx":
        return any(n.startswith("word/media/") for n in zipfile.ZipFile(yol).namelist())
    if uzanti == "odt":  # odfpy/LibreOffice "Pictures/", Word "media/" kullanır
        return any(n.startswith(("Pictures/", "media/"))
                   for n in zipfile.ZipFile(yol).namelist())
    if uzanti == "html":
        return "<img" in yol.read_text(encoding="utf-8", errors="replace").lower()
    if uzanti == "md":
        return "![" in yol.read_text(encoding="utf-8")
    if uzanti == "pdf":
        from pypdf import PdfReader

        return any(len(s.images) for s in PdfReader(str(yol)).pages)
    raise ValueError(uzanti)


# ============================================================ ses/video örnekleri

def ffmpeg(*argumanlar):
    import imageio_ffmpeg

    sonuc = subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-hide_banner", "-y",
                            *map(str, argumanlar)], capture_output=True, text=True,
                           encoding="utf-8", errors="replace")
    assert sonuc.returncode == 0, sonuc.stderr[-800:]
    return sonuc


def okunabilir_medya(yol: Path) -> bool:
    """ffmpeg dosyayı hatasız baştan sona çözebiliyor mu?"""
    import imageio_ffmpeg

    sonuc = subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-v", "error", "-i", str(yol),
                            "-f", "null", "-"], capture_output=True, text=True,
                           encoding="utf-8", errors="replace")
    return sonuc.returncode == 0 and yol.stat().st_size > 0
