"""Belge blok modeli: biçimlendirme koruyan ara temsil.

Bir belge Blok listesine okunur; her blok başlık/paragraf/liste maddesi/resim
olabilir ve metin parçaları kalın/italik bilgisi taşır. Her hedef format bu
blokları kendi diline çevirir. Böylece yazılar, başlıklar, kalın/italik
vurgular ve gömülü resimler dönüşümde korunur.
"""

import base64
import binascii
import html
import html.parser
import io
import re
import sys
import urllib.parse
import urllib.request
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

# Blok tipleri: "p", "h1".."h6", "li", "resim"

# XML 1.0'da geçersiz karakterler: DOCX yazıcısı bunlarla çöküyor, ODT
# yazıcısı ise sessizce açılamayan bozuk dosya üretiyordu (ör. PDF'ten
# çıkarılan metinlerdeki \x00, TXT'lerdeki kontrol karakterleri). Sayfa
# sonu / dikey sekme boşluğa çevrilir, geri kalanlar atılır.
_BOSLUGA = str.maketrans({"\x0b": " ", "\x0c": " "})
_XML_GECERSIZ = re.compile("[\x00-\x08\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def xml_temizle(metin: str) -> str:
    return _XML_GECERSIZ.sub("", metin.translate(_BOSLUGA))


@dataclass
class Parca:
    metin: str
    kalin: bool = False
    italik: bool = False

    def __post_init__(self):
        self.metin = xml_temizle(self.metin)


@dataclass
class Blok:
    tip: str = "p"
    parcalar: list[Parca] = field(default_factory=list)
    resim: bytes | None = None
    resim_ad: str = ""

    def metin(self) -> str:
        return "".join(p.metin for p in self.parcalar)


def _bos_mu(blok: Blok) -> bool:
    return blok.tip == "p" and not blok.resim and not blok.metin().strip()


_ARIAL = r"C:\Windows\Fonts\arial.ttf"
_ARIAL_KALIN = r"C:\Windows\Fonts\arialbd.ttf"
_ARIAL_ITALIK = r"C:\Windows\Fonts\ariali.ttf"
_ARIAL_KI = r"C:\Windows\Fonts\arialbi.ttf"

# Windows'un ANSI kod sayfası (Türkçe sistemlerde cp1254)
_ANSI = "mbcs" if sys.platform == "win32" else "cp1254"


def metin_coz(veri: bytes, html_mi: bool = False) -> str:
    """Bayt içeriğini doğru kodlamayla metne çevirir.

    Sıra: BOM -> (HTML ise <meta charset>) -> UTF-8 -> Windows ANSI.
    Eskiden her şey UTF-8 sanılıyordu; Not Defteri'nin eski sürümleriyle ya
    da Word'le "ANSI" kaydedilmiş Türkçe dosyalarda ş/ğ/ı gibi harfler "�"
    oluyordu.
    """
    if veri.startswith(b"\xef\xbb\xbf"):
        return veri[3:].decode("utf-8", errors="replace")
    if veri.startswith((b"\xff\xfe", b"\xfe\xff")):
        return veri.decode("utf-16", errors="replace")
    if html_mi:
        eslesme = re.search(
            rb"""<meta[^>]+charset\s*=\s*["']?\s*([A-Za-z0-9_\-]+)""", veri[:4096], re.I)
        if eslesme:
            try:
                return veri.decode(eslesme.group(1).decode("ascii"))
            except (LookupError, UnicodeDecodeError):
                pass  # bildirilen kodlama yanlış/bilinmiyor -- aşağıdaki sıraya düş
    try:
        return veri.decode("utf-8")
    except UnicodeDecodeError:
        return veri.decode(_ANSI, errors="replace")


# ============================================================ RESİM GÜVENLİĞİ

_AZAMI_RESIM_BAYT = 50 * 1024 * 1024

_MIME_UZANTI = {
    "image/png": ".png", "image/jpeg": ".jpg", "image/jpg": ".jpg",
    "image/gif": ".gif", "image/webp": ".webp", "image/bmp": ".bmp",
    "image/svg+xml": ".svg", "image/tiff": ".tif", "image/x-icon": ".ico",
    "image/vnd.microsoft.icon": ".ico", "image/avif": ".avif",
}
_RESIM_UZANTILARI = {
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg", ".tif", ".tiff",
    ".ico", ".avif", ".heic", ".emf", ".wmf", ".jfif",
}
_WINDOWS_AYRILMIS = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(10)),
                     *(f"LPT{i}" for i in range(10))}


def _resim_mi(yol: Path) -> bool:
    try:
        if yol.suffix.lower() == ".svg":
            with open(yol, "rb") as f:
                return b"<svg" in f.read(4096).lower()
        from PIL import Image

        with Image.open(yol):
            return True
    except Exception:
        return False


def yerel_resim_yolu(adres: str, taban: Path | None) -> Path | None:
    """HTML/MD içindeki bir resim adresini güvenli biçimde yerel bir resim
    dosyasına çözer; uygun değilse None döndürür.

    Güvenlik: Eskiden adres hiç denetlenmiyordu. Güvenilmeyen bir HTML/MD
    dosyası "C:/Users/.../gizli.txt" gibi bir adresle bilgisayardaki herhangi
    bir dosyanın içeriğini dönüşüm çıktısına gömdürebiliyordu; "\\\\sunucu\\x.png"
    gibi bir ağ (UNC) yolu ise Windows'un oturum kimlik bilgisi özetini
    (NTLM) o sunucuya göndermesine yol açabiliyordu. Artık ağ yolları ve URL
    şemaları reddedilir, yalnızca gerçekten resim olan yerel dosyalar kabul
    edilir.
    """
    if taban is None:
        return None
    adres = (adres or "").strip()
    if not adres:
        return None
    if adres.lower().startswith("file:"):
        ayrik = urllib.parse.urlsplit(adres)
        if ayrik.netloc.lower() not in ("", "localhost"):
            return None  # file://sunucu/... = ağ yolu
        adres = urllib.request.url2pathname(ayrik.path)
    elif re.match(r"^[a-zA-Z][a-zA-Z0-9+.\-]*:", adres) and not re.match(
            r"^[a-zA-Z]:([\\/]|$)", adres):
        return None  # http:, https:, data:, javascript: vb. (sürücü harfi hariç)
    else:
        adres = urllib.parse.unquote(adres.split("#")[0].split("?")[0])
    if adres.replace("/", "\\").startswith("\\\\"):
        return None  # UNC ağ yolu (\\sunucu\paylasim, //sunucu, \\?\UNC\...)

    yol = Path(adres)
    if not yol.is_absolute():
        yol = taban / yol
    try:
        yol = yol.resolve(strict=True)
        if not yol.is_file() or yol.stat().st_size > _AZAMI_RESIM_BAYT:
            return None
    except (OSError, RuntimeError, ValueError):
        return None
    return yol if _resim_mi(yol) else None


def data_uri_coz(adres: str) -> tuple[bytes, str] | None:
    """data:image/...;base64,... adresini (veri, güvenli_ad) olarak çözer."""
    bas, virgul, govde = adres.partition(",")
    if not virgul or not bas.lower().startswith("data:"):
        return None
    ozellikler = [o.strip().lower() for o in bas[5:].split(";")]
    uzanti = _MIME_UZANTI.get(ozellikler[0])
    if not uzanti:
        return None
    try:
        if "base64" in ozellikler[1:]:
            veri = base64.b64decode(re.sub(r"\s+", "", govde))
        else:
            veri = urllib.parse.unquote_to_bytes(govde)
    except (ValueError, binascii.Error):
        return None
    if not veri or len(veri) > _AZAMI_RESIM_BAYT:
        return None
    # Ad MIME türünden DEĞİL sabit tablodan üretilir: eskiden
    # "data:image/..\..\..\x.bat;base64,..." gibi bir adres, HTML -> MD
    # dönüşümünde istenen klasöre dosya yazdırabiliyordu (yol geçişi).
    return veri, "resim" + uzanti


def guvenli_dosya_adi(ad: str, varsayilan: str = "resim.png") -> str:
    """Belgeden gelen bir resim adını tek parçalı, zararsız bir dosya adına
    çevirir (klasör ayırıcıları, '..', Windows'ta yasak karakterler ve
    resim olmayan uzantılar temizlenir)."""
    ad = re.split(r"[\\/]", ad or "")[-1]
    ad = re.sub(r'[<>:"|?*\x00-\x1f]', "_", ad).strip(" .")[:100]
    if not ad:
        return varsayilan
    kok, nokta, uzanti = ad.rpartition(".")
    if not nokta:
        kok, uzanti = ad, ""
    if "." + uzanti.lower() not in _RESIM_UZANTILARI:
        kok, uzanti = ad, "bin"
    if kok.upper() in _WINDOWS_AYRILMIS or not kok:
        kok = "_" + kok
    return f"{kok}.{uzanti}"


_IMG_ETIKETI = re.compile(r"<img\b[^>]*>", re.I)
_SRC = re.compile(r"""(\bsrc\s*=\s*)("[^"]*"|'[^']*'|[^\s>]+)""", re.I)


def resim_adreslerini_yerellestir(html_metin: str, taban: Path, gecici: Path) -> str:
    """Markdown'dan üretilen HTML'deki resim adreslerini Word'ün açabileceği
    mutlak dosya yollarına çevirir (Word ara HTML'i geçici klasörde açtığı
    için göreli adresli resimler kayboluyordu). data: resimleri geçici
    dosyaya yazılır; güvenli olmayan (ağ yolu, URL) ya da bulunamayan
    resimlerin etiketi tamamen kaldırılır (boş src bırakılınca Word sayfanın
    kendisini resim diye gömüyordu).

    NOT: Ham Windows yolu kullanılır; Word yüzde-kodlu ("%C3%B6") file://
    adreslerindeki Türkçe karakterleri çözemiyor."""
    sayac = [0]

    def degistir(etiket):
        src = _SRC.search(etiket.group(0))
        if not src:
            return ""
        adres = html.unescape(src.group(2).strip("\"'"))
        yeni = None
        if adres.lower().startswith("data:"):
            cozum = data_uri_coz(adres)
            if cozum:
                sayac[0] += 1
                dosya = gecici / f"gomulu_{sayac[0]}{Path(cozum[1]).suffix}"
                dosya.write_bytes(cozum[0])
                yeni = str(dosya)
        else:
            yol = yerel_resim_yolu(adres, taban)
            if yol is not None:
                yeni = str(yol)
        if yeni is None:
            return ""
        return (etiket.group(0)[:src.start()] + f'{src.group(1)}"{html.escape(yeni)}"'
                + etiket.group(0)[src.end():])

    return _IMG_ETIKETI.sub(degistir, html_metin)


# Açılışta dışarıdan kaynak yükletebilen adresler (resim, stil, çerçeve, CSS url())
_KAYNAK_ADRESI = re.compile(
    r"""\b(?:src|href|background|data|poster|lowsrc|dynsrc|codebase)\s*=\s*"""
    r"""("[^"]*"|'[^']*'|[^\s>]+)|url\(\s*("[^"]*"|'[^']*'|[^)\s]+)"""
    r"""|@import\s+("[^"]*"|'[^']*')""", re.I)


def ag_adresi_var(html_metin: str) -> bool:
    """HTML ağ paylaşımı (UNC) ya da uzak file:// adresine başvuruyor mu?

    Word bir HTML'i açarken bu adreslere bağlanır; "\\\\sunucu\\x.png" gibi bir
    adres Windows'un oturum kimlik bilgisi özetini (NTLM) o sunucuya
    gönderebilir. Böyle belgeler Word'e verilmez, ağa hiç çıkmayan dahili
    motorla dönüştürülür."""
    for eslesme in _KAYNAK_ADRESI.finditer(html_metin):
        deger = html.unescape(next(g for g in eslesme.groups() if g is not None)).strip("\"' ")
        if deger.replace("/", "\\").startswith("\\\\"):
            return True  # \\sunucu\pay, //sunucu/pay (Word bunu UNC sayabilir)
        if deger.lower().startswith("file:"):
            if urllib.parse.urlsplit(deger).netloc.lower() not in ("", "localhost"):
                return True
            yol = urllib.request.url2pathname(urllib.parse.urlsplit(deger).path)
            if yol.replace("/", "\\").startswith("\\\\"):
                return True
    return False


def _resim_olcusu(veri: bytes) -> tuple[int, int] | None:
    try:
        from PIL import Image

        with Image.open(io.BytesIO(veri)) as im:
            return im.size
    except Exception:
        return None


def _png_e_cevir(veri: bytes) -> bytes | None:
    """Hedefin doğrudan desteklemediği resimleri (WEBP, EMF, CMYK JPEG,
    16-bit PNG...) standart 8-bit PNG'ye çevirir."""
    try:
        from PIL import Image

        from .gorsel import _cmyk_rgb, _sekiz_bite_indir

        with Image.open(io.BytesIO(veri)) as im:
            im.load()
            kare = _sekiz_bite_indir(im)
            if kare.mode == "CMYK":
                kare = _cmyk_rgb(kare)
            kare = kare.convert("RGBA" if kare.has_transparency_data else "RGB")
        cikti = io.BytesIO()
        kare.save(cikti, "PNG")
        return cikti.getvalue()
    except Exception:
        return None


def _pdf_icin_resim(veri: bytes) -> tuple[bytes, tuple[int, int]] | None:
    """reportlab'in güvenle gömebileceği (8-bit RGB/gri JPEG ya da PNG)
    resim verisini ve piksel ölçüsünü döndürür; gerekirse PNG'ye çevirir."""
    try:
        from PIL import Image

        with Image.open(io.BytesIO(veri)) as im:
            uygun = ((im.format == "JPEG" and im.mode in ("RGB", "L"))
                     or (im.format == "PNG" and im.mode in ("RGB", "RGBA", "L", "LA", "P")))
            if uygun:
                return veri, im.size
    except Exception:
        return None
    png = _png_e_cevir(veri)
    olcu = _resim_olcusu(png) if png else None
    return (png, olcu) if olcu else None


# ================================================================= OKUYUCULAR

def _satir_bloklari(metin: str) -> list[Blok]:
    return [Blok("p", [Parca(satir.rstrip())]) for satir in metin.splitlines()]


def oku_txt(yol: Path) -> list[Blok]:
    return _satir_bloklari(metin_coz(yol.read_bytes()))


def oku_rtf(yol: Path) -> list[Blok]:
    from striprtf.striprtf import rtf_to_text

    return _satir_bloklari(rtf_to_text(metin_coz(yol.read_bytes()), errors="replace"))


def oku_pdf(yol: Path) -> list[Blok]:
    from pypdf import PdfReader
    from pypdf.errors import FileNotDecryptedError

    okuyucu = PdfReader(str(yol))
    if okuyucu.is_encrypted:
        try:
            acildi = okuyucu.decrypt("")  # çoğu "sahip parolalı" PDF boş parolayla açılır
        except Exception:
            acildi = 0
        if not acildi:
            raise RuntimeError(
                "Bu PDF parola ile korunuyor; parola olmadan dönüştürülemiyor.")

    bloklar: list[Blok] = []
    try:
        sayfalar = list(okuyucu.pages)
    except FileNotDecryptedError:
        raise RuntimeError(
            "Bu PDF parola ile korunuyor; parola olmadan dönüştürülemiyor.") from None
    for sayfa in sayfalar:
        metin = sayfa.extract_text() or ""
        for satir in metin.splitlines():
            bloklar.append(Blok("p", [Parca(satir.rstrip())]))
        bloklar.append(Blok("p", [Parca("")]))
    while bloklar and not bloklar[-1].metin():
        bloklar.pop()
    return bloklar


_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_OLE_IMZASI = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def oku_docx(yol: Path) -> list[Blok]:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run

    with open(yol, "rb") as f:
        if f.read(8) == _OLE_IMZASI:
            # Parolalı DOCX'ler (ve .docx diye adlandırılmış eski .doc'lar)
            # ZIP değil OLE kabıdır; python-docx anlaşılmaz bir hata veriyordu.
            raise RuntimeError(
                "Bu belge parola ile korunuyor (ya da eski .doc biçiminde); "
                "bu haliyle dönüştürülemiyor.")

    belge = docx.Document(str(yol))
    bloklar: list[Blok] = []
    # Paragrafın kendi metni sayılmayan run'lar: metin kutuları, silinmiş /
    # taşınmış (izlenen değişiklik) metin, uyumluluk yedekleri (yinelenen içerik)
    atlanacak = {qn("w:txbxContent"), qn("w:del"), qn("w:moveFrom"), _MC_FALLBACK}

    def runlar(p):
        # p.runs yalnızca doğrudan w:r çocuklarını verir; köprü (w:hyperlink),
        # içerik denetimi (w:sdt), alan (w:fldSimple) ve izlenen ekleme (w:ins)
        # içindeki metin eskiden tamamen kayboluyordu.
        for r in p._p.iter(qn("w:r")):
            ata = r.getparent()
            while ata is not None and ata is not p._p:
                if ata.tag in atlanacak:
                    break
                ata = ata.getparent()
            else:
                yield Run(r, p)

    def paragrafi_isle(p):
        stil = ((p.style.name if p.style is not None else "") or "").lower()
        eslesme = re.match(r"heading (\d)", stil)
        if eslesme:
            tip = "h" + eslesme.group(1)
        elif "list" in stil:
            tip = "li"
        else:
            tip = "p"
        blok = Blok(tip)
        for run in runlar(p):
            # Gömülü resimleri ayıkla
            for blip in run._element.findall(".//" + qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                if rid and rid in p.part.related_parts:
                    parca = p.part.related_parts[rid]
                    bloklar.append(Blok(
                        "resim", resim=parca.blob,
                        resim_ad=Path(parca.partname).name))
            if run.text:
                blok.parcalar.append(Parca(
                    run.text, bool(run.bold), bool(run.italic)))
        if not blok.parcalar and p.text:
            blok.parcalar.append(Parca(p.text))
        bloklar.append(blok)

    def tabloyu_isle(tablo):
        for satir in tablo.rows:
            try:
                hucreler, onceki = [], None
                for hucre in satir.cells:
                    if hucre._tc is onceki:  # yatay birleştirilmiş hücre tekrarı
                        continue
                    onceki = hucre._tc
                    hucreler.append(" ".join(hucre.text.split()))
                metin = " | ".join(hucreler)
            except Exception:  # bozuk tablo yapısı: en azından metni kurtar
                metin = " ".join("".join(satir._tr.itertext()).split())
            bloklar.append(Blok("p", [Parca(metin)]))

    def gez(kap):
        for eleman in kap.iterchildren():
            if eleman.tag == qn("w:p"):
                paragrafi_isle(Paragraph(eleman, belge))
            elif eleman.tag == qn("w:tbl"):
                tabloyu_isle(Table(eleman, belge))
            elif eleman.tag == qn("w:sdt"):
                # Blok düzeyi içerik denetimleri (kapak sayfası, içindekiler...)
                # eskiden tamamen atlanıyordu.
                icerik = eleman.find(qn("w:sdtContent"))
                if icerik is not None:
                    gez(icerik)
            elif eleman.tag == qn("w:customXml"):
                gez(eleman)

    gez(belge.element.body)
    return bloklar


class _HTMLBlokAyiklayici(html.parser.HTMLParser):
    _BASLIKLAR = {"h1", "h2", "h3", "h4", "h5", "h6"}
    _ATLA = {"script", "style", "head", "title", "template"}
    _BLOK = {"p", "div", "tr", "blockquote", "pre", "table", "ul", "ol", "dl",
             "dt", "dd", "section", "article", "header", "footer", "figure",
             "figcaption", "main", "nav", "aside", "address", "hr"}

    def __init__(self, taban_klasor: Path | None = None):
        super().__init__()
        self.bloklar: list[Blok] = []
        self.taban = taban_klasor
        self._tip = "p"
        self._parcalar: list[Parca] = []
        self._kalin = 0
        self._italik = 0
        self._atla = 0
        self._pre = 0
        self._pre_basi = False

    def _icerik_var(self) -> bool:
        return any(p.metin.strip() for p in self._parcalar)

    def _bitir(self):
        if self._icerik_var():
            self.bloklar.append(Blok(self._tip, self._parcalar))
        elif self.bloklar and self.bloklar[-1].metin().strip():
            self.bloklar.append(Blok("p", [Parca("")]))
        self._parcalar = []
        self._tip = "p"

    def _satir_sonu(self):
        """<br> ya da <pre> içindeki satır sonu: bloğu bitir, tipi koru."""
        tip = self._tip
        if self._icerik_var():
            self._bitir()
        else:
            self._parcalar = []
            self.bloklar.append(Blok("p", [Parca("")]))
        self._tip = tip

    def handle_starttag(self, tag, attrs):
        if tag in self._ATLA:
            self._atla += 1
        elif tag in self._BASLIKLAR:
            self._bitir()
            self._tip = tag
        elif tag == "li":
            self._bitir()
            self._tip = "li"
        elif tag == "br":
            self._satir_sonu()
        elif tag in self._BLOK:
            # <li><p>metin</p></li>: madde tipi korunur (eskiden <p> tipi
            # "p"ye sıfırladığı için madde işaretleri kayboluyordu)
            devam = "p" if self._icerik_var() else self._tip
            self._bitir()
            self._tip = devam
            if tag == "pre":
                self._pre += 1
                self._pre_basi = True
        elif tag in ("td", "th"):
            # Hücreleri ayır (eskiden "<td>a</td><td>b</td>" -> "ab" oluyordu)
            if self._icerik_var():
                self._parcalar.append(Parca(" | "))
        elif tag in ("b", "strong"):
            self._kalin += 1
        elif tag in ("i", "em"):
            self._italik += 1
        elif tag == "img":
            self._resim(dict(attrs).get("src") or "")

    def handle_endtag(self, tag):
        if tag in self._ATLA:
            self._atla = max(0, self._atla - 1)
        elif tag in self._BASLIKLAR or tag == "li" or tag in self._BLOK:
            self._bitir()
            if tag == "pre":
                self._pre = max(0, self._pre - 1)
        elif tag in ("b", "strong"):
            self._kalin = max(0, self._kalin - 1)
        elif tag in ("i", "em"):
            self._italik = max(0, self._italik - 1)

    def handle_data(self, data):
        if self._atla or not data:
            return
        kalin, italik = self._kalin > 0, self._italik > 0
        if self._pre:
            # Kod blokları: satır sonları ve girintiler korunur
            if self._pre_basi and data.startswith("\n"):
                data = data[1:]
            self._pre_basi = False
            for i, satir in enumerate(data.split("\n")):
                if i:
                    self._satir_sonu()
                if satir:
                    self._parcalar.append(Parca(satir, kalin, italik))
        else:
            self._parcalar.append(Parca(re.sub(r"\s+", " ", data), kalin, italik))

    def _resim(self, src):
        src = src.strip()
        if src.lower().startswith("data:"):
            cozum = data_uri_coz(src)
            if not cozum:
                return
            veri, ad = cozum
        else:
            yol = yerel_resim_yolu(src, self.taban)
            if yol is None:
                return
            try:
                veri = yol.read_bytes()
            except OSError:
                return
            ad = yol.name
        self._bitir()
        self.bloklar.append(Blok("resim", resim=veri, resim_ad=ad))


def _bosluklari_duzenle(bloklar: list[Blok]) -> list[Blok]:
    """Ayraç boş satırlarını toparlar: baştaki/sondaki ve ardışık boşlar
    atılır; liste maddeleri arasına boş satır girmez (girince HTML/DOCX'te
    liste her maddede bölünüyordu)."""
    sonuc: list[Blok] = []
    for i, blok in enumerate(bloklar):
        if _bos_mu(blok):
            if not sonuc or _bos_mu(sonuc[-1]):
                continue
            sonraki = bloklar[i + 1] if i + 1 < len(bloklar) else None
            if sonuc[-1].tip == "li" and sonraki is not None and sonraki.tip == "li":
                continue
        sonuc.append(blok)
    while sonuc and _bos_mu(sonuc[-1]):
        sonuc.pop()
    return sonuc


def oku_html_metin(icerik: str, taban: Path | None = None) -> list[Blok]:
    ayiklayici = _HTMLBlokAyiklayici(taban)
    ayiklayici.feed(icerik)
    ayiklayici.close()
    ayiklayici._bitir()
    return _bosluklari_duzenle(ayiklayici.bloklar)


def oku_html(yol: Path) -> list[Blok]:
    return oku_html_metin(metin_coz(yol.read_bytes(), html_mi=True), yol.parent)


def md_html(metin: str) -> str:
    import markdown

    return markdown.markdown(metin, extensions=["tables", "fenced_code"])


def oku_md(yol: Path) -> list[Blok]:
    return oku_html_metin(md_html(metin_coz(yol.read_bytes())), yol.parent)


# ODT'de içine girilip okunan kapsayıcılar (bölümler, içindekiler gövdesi...).
# "*-source" (içindekiler şablonu) ve izlenen değişiklikler bilerek dışarıda.
_ODT_KAPLAR = {"section", "index-body", "index-title", "table-of-content",
               "illustration-index", "table-index", "object-index", "user-index",
               "alphabetical-index", "bibliography"}


def oku_odt(yol: Path) -> list[Blok]:
    from odf import teletype
    from odf.draw import Image as DrawImage
    from odf.namespaces import DRAWNS, FONS, OFFICENS, STYLENS, TABLENS, TEXTNS, XLINKNS
    from odf.opendocument import load

    belge = load(str(yol))
    bloklar: list[Blok] = []

    # ---- stil adı -> (kalın, italik), üst stillerden kalıtım dahil
    ham_stiller: dict[str, tuple] = {}
    for kap in (belge.styles, belge.automaticstyles):
        for stil in kap.childNodes:
            if getattr(stil, "qname", None) != (STYLENS, "style"):
                continue
            kalin = italik = None
            for ozellik in stil.childNodes:
                if getattr(ozellik, "qname", None) == (STYLENS, "text-properties"):
                    agirlik = ozellik.getAttrNS(FONS, "font-weight")
                    egim = ozellik.getAttrNS(FONS, "font-style")
                    if agirlik:
                        kalin = agirlik == "bold" or (agirlik.isdigit() and int(agirlik) >= 600)
                    if egim:
                        italik = egim in ("italic", "oblique")
            ham_stiller[stil.getAttrNS(STYLENS, "name")] = (
                kalin, italik, stil.getAttrNS(STYLENS, "parent-style-name"))

    def stil_coz(ad, derinlik=0) -> tuple[bool, bool]:
        if not ad or ad not in ham_stiller or derinlik > 10:
            return False, False
        kalin, italik, ust = ham_stiller[ad]
        ust_kalin, ust_italik = stil_coz(ust, derinlik + 1)
        return (ust_kalin if kalin is None else kalin,
                ust_italik if italik is None else italik)

    # NOT: odf.text.H()/.P()/.List() gibi fabrika fonksiyonları, sadece .qname
    # okumak için bile çağrılsalar zorunlu alan denetimi yapıp hata fırlatır
    # (ör. H() -> "outlinelevel" eksik). Bu yüzden qname'ler doğrudan sabit
    # (TEXTNS, etiket) demetiyle karşılaştırılır; hiçbir eleman örneği kurulmaz.
    def parcalar(eleman, kalin=False, italik=False) -> list[Parca]:
        sonuc: list[Parca] = []
        for dugum in eleman.childNodes:
            q = getattr(dugum, "qname", None)
            if q is None:
                if dugum.nodeType == dugum.TEXT_NODE:
                    sonuc.append(Parca(str(dugum.data), kalin, italik))
            elif q == (TEXTNS, "span"):
                k, i = stil_coz(dugum.getAttrNS(TEXTNS, "style-name"))
                sonuc += parcalar(dugum, kalin or k, italik or i)
            elif q == (TEXTNS, "s"):
                sonuc.append(Parca(" " * int(dugum.getAttrNS(TEXTNS, "c") or 1), kalin, italik))
            elif q == (TEXTNS, "tab"):
                sonuc.append(Parca("\t", kalin, italik))
            elif q == (TEXTNS, "line-break"):
                sonuc.append(Parca(" ", kalin, italik))
            elif q in ((TEXTNS, "note"), (TEXTNS, "tracked-changes")) or q[0] in (
                    DRAWNS, OFFICENS):
                continue  # dipnot gövdesi, yorum, çizim/metin kutusu
            else:
                sonuc += parcalar(dugum, kalin, italik)  # köprü, alan vb.
        return sonuc

    # odfpy yalnızca "Pictures/" altındaki resimleri yükler; Word gibi
    # programlar ODT'de resimleri "media/" altına koyar -- bunlar arşivden
    # doğrudan okunur. (Dış dosyaya bağlantılı resimler bilerek okunmaz.)
    with zipfile.ZipFile(yol) as arsiv:
        arsiv_adlari = set(arsiv.namelist())

    def resimleri_ekle(eleman):
        for resim in eleman.getElementsByType(DrawImage):
            href = resim.getAttrNS(XLINKNS, "href") or ""
            kayit = belge.Pictures.get(href)
            veri = kayit[1] if kayit else None
            if not veri and href in arsiv_adlari:
                with zipfile.ZipFile(yol) as arsiv:
                    veri = arsiv.read(href)
            if veri and len(veri) <= _AZAMI_RESIM_BAYT:
                bloklar.append(Blok("resim", resim=veri, resim_ad=Path(href).name))

    def tablo(eleman):
        def satirlar(kap):
            for c in kap.childNodes:
                q = getattr(c, "qname", None)
                if q == (TABLENS, "table-row"):
                    yield c
                elif q in ((TABLENS, "table-header-rows"), (TABLENS, "table-rows"),
                           (TABLENS, "table-row-group")):
                    yield from satirlar(c)

        for satir in satirlar(eleman):
            hucreler = [
                " ".join(" ".join(teletype.extractText(c) for c in hucre.childNodes
                                  if getattr(c, "qname", None)).split())
                for hucre in satir.childNodes
                if getattr(hucre, "qname", None) == (TABLENS, "table-cell")]
            while hucreler and not hucreler[-1]:
                hucreler.pop()
            bloklar.append(Blok("p", [Parca(" | ".join(hucreler))]))

    def gez(dugum, liste_icinde=False):
        for eleman in dugum.childNodes:
            q = getattr(eleman, "qname", None)
            if q is None:
                continue
            if q == (TEXTNS, "h"):
                resimleri_ekle(eleman)
                seviye = int(eleman.getAttrNS(TEXTNS, "outline-level") or 1)
                bloklar.append(Blok(f"h{min(max(seviye, 1), 6)}", parcalar(eleman)))
            elif q == (TEXTNS, "p"):
                resimleri_ekle(eleman)
                k, i = stil_coz(eleman.getAttrNS(TEXTNS, "style-name"))
                bloklar.append(Blok("li" if liste_icinde else "p", parcalar(eleman, k, i)))
            elif q == (TEXTNS, "list"):
                gez(eleman, True)
            elif q in ((TEXTNS, "list-item"), (TEXTNS, "list-header")):
                gez(eleman, liste_icinde)
            elif q == (TABLENS, "table"):
                tablo(eleman)
            elif q[0] == TEXTNS and q[1] in _ODT_KAPLAR:
                gez(eleman, liste_icinde)

    gez(belge.text)
    return bloklar


OKUYUCULAR = {
    "txt": oku_txt,
    "md": oku_md,
    "html": oku_html,
    "docx": oku_docx,
    "pdf": oku_pdf,
    "rtf": oku_rtf,
    "odt": oku_odt,
}


# ================================================================== YAZICILAR

def yaz_txt(bloklar: list[Blok], yol: Path):
    satirlar = []
    for blok in bloklar:
        if blok.tip == "resim":
            satirlar.append(f"[Resim: {guvenli_dosya_adi(blok.resim_ad)}]")
        elif blok.tip == "li":
            satirlar.append("• " + blok.metin())
        else:
            satirlar.append(blok.metin())
    yol.write_text("\n".join(satirlar), encoding="utf-8")


# Markdown'da anlamı olan karakterler kaçırılır; yoksa düz metindeki "*",
# "[...]" ya da "<script>" gibi ifadeler MD görüntüleyicide biçim/HTML olarak
# yorumlanıyordu.
_MD_OZEL = re.compile(r"([\\`*\[\]<>])")
_MD_SATIR_BASI = re.compile(r"^(\s*)([#>+\-=|]|\d+[.)](?=\s|$))")


def _md_kacir(metin: str) -> str:
    metin = _MD_OZEL.sub(r"\\\1", metin)
    return re.sub(r"(?<!\w)_|_(?!\w)", r"\\_", metin)


def _md_satir_basi_kacir(metin: str) -> str:
    eslesme = _MD_SATIR_BASI.match(metin)
    if not eslesme:
        return metin
    girinti, isaret = eslesme.groups()
    if isaret[0].isdigit():  # "1. madde" -> "1\. madde"
        kacmis = isaret[:-1] + "\\" + isaret[-1]
    else:
        kacmis = "\\" + isaret
    return girinti + kacmis + metin[eslesme.end():]


def _benzersiz_ad(klasor: Path, ad: str, kullanilan: set[str]) -> str:
    kok, uzanti = Path(ad).stem, Path(ad).suffix
    aday, sira = ad, 2
    while aday.lower() in kullanilan or (klasor / aday).exists():
        aday = f"{kok}_{sira}{uzanti}"
        sira += 1
    kullanilan.add(aday.lower())
    return aday


def yaz_md(bloklar: list[Blok], yol: Path):
    parcalar_md: list[tuple[str, str]] = []  # (tip, satır)
    resim_klasoru = yol.parent / (yol.stem + "_dosyalar")
    kullanilan: set[str] = set()
    for blok in bloklar:
        if blok.tip == "resim":
            resim_klasoru.mkdir(exist_ok=True)
            # Ad güvenli tek parçaya indirgenir ve benzersizleştirilir: eskiden
            # aynı adlı ("resim.png") tüm resimler birbirinin üzerine yazılıyordu.
            ad = _benzersiz_ad(resim_klasoru, guvenli_dosya_adi(blok.resim_ad), kullanilan)
            (resim_klasoru / ad).write_bytes(blok.resim)
            # Adres URL-kodlanır: klasör adındaki boşluklar bağlantıyı kırıyordu
            adres = urllib.parse.quote(f"{resim_klasoru.name}/{ad}")
            parcalar_md.append(("resim", f"![{_md_kacir(Path(ad).stem)}]({adres})"))
            continue
        metin = ""
        for parca in blok.parcalar:
            m = parca.metin
            if not m.strip():
                metin += m
                continue
            ic = _md_kacir(m.strip())
            if parca.kalin and parca.italik:
                ic = f"***{ic}***"
            elif parca.kalin:
                ic = f"**{ic}**"
            elif parca.italik:
                ic = f"*{ic}*"
            bas = " " if m[:1].isspace() else ""
            son = " " if m[-1:].isspace() else ""
            metin += bas + ic + son
        metin = metin.strip()
        if blok.tip.startswith("h"):
            parcalar_md.append((blok.tip, "#" * int(blok.tip[1]) + " " + metin))
        elif blok.tip == "li":
            parcalar_md.append(("li", "- " + _md_satir_basi_kacir(metin)))
        elif metin:
            parcalar_md.append(("p", _md_satir_basi_kacir(metin)))

    cikti = ""
    for i, (tip, satir) in enumerate(parcalar_md):
        if i:
            # ardışık liste maddeleri tek satır sonuyla (sıkı liste)
            cikti += "\n" if tip == "li" and parcalar_md[i - 1][0] == "li" else "\n\n"
        cikti += satir
    yol.write_text(cikti + "\n", encoding="utf-8")


def yaz_html(bloklar: list[Blok], yol: Path):
    import mimetypes

    parcalar_html: list[str] = []
    liste_acik = False
    for blok in bloklar:
        if blok.tip == "li" and not liste_acik:
            parcalar_html.append("<ul>")
            liste_acik = True
        elif blok.tip != "li" and liste_acik:
            parcalar_html.append("</ul>")
            liste_acik = False

        if blok.tip == "resim":
            ad = guvenli_dosya_adi(blok.resim_ad)
            tur = mimetypes.guess_type(ad)[0] or "image/png"
            veri = base64.b64encode(blok.resim).decode()
            parcalar_html.append(
                f'<p><img src="data:{tur};base64,{veri}" '
                f'alt="{html.escape(ad)}" style="max-width:100%"></p>')
            continue

        ic = ""
        for parca in blok.parcalar:
            m = html.escape(parca.metin)
            if parca.kalin:
                m = f"<strong>{m}</strong>"
            if parca.italik:
                m = f"<em>{m}</em>"
            ic += m
        if blok.tip.startswith("h"):
            parcalar_html.append(f"<{blok.tip}>{ic}</{blok.tip}>")
        elif blok.tip == "li":
            parcalar_html.append(f"<li>{ic}</li>")
        elif ic.strip():
            parcalar_html.append(f"<p>{ic}</p>")
        else:
            parcalar_html.append("<br>")
    if liste_acik:
        parcalar_html.append("</ul>")

    sablon = (
        "<!DOCTYPE html>\n"
        '<html lang="tr">\n<head>\n<meta charset="utf-8">\n'
        f"<title>{html.escape(yol.stem)}</title>\n"
        "<style>body{font-family:Arial,sans-serif;max-width:50rem;"
        "margin:2rem auto;padding:0 1rem;line-height:1.6;}</style>\n"
        "</head>\n<body>\n" + "\n".join(parcalar_html) + "\n</body>\n</html>\n"
    )
    yol.write_text(sablon, encoding="utf-8")


def yaz_docx(bloklar: list[Blok], yol: Path):
    import docx

    belge = docx.Document()
    for blok in bloklar:
        if blok.tip == "resim":
            if not _docx_resim_ekle(belge, blok.resim):
                belge.add_paragraph(f"[Resim: {guvenli_dosya_adi(blok.resim_ad)}]")
            continue
        if blok.tip.startswith("h"):
            p = belge.add_heading(level=int(blok.tip[1]))
        elif blok.tip == "li":
            p = belge.add_paragraph(style="List Bullet")
        else:
            p = belge.add_paragraph()
        for parca in blok.parcalar:
            run = p.add_run(parca.metin)
            run.bold = parca.kalin
            run.italic = parca.italik
    belge.save(str(yol))


def _docx_resim_ekle(belge, veri: bytes) -> bool:
    """Resmi gerçek boyutunda (96 DPI), sayfaya sığacak şekilde ekler.
    Eskiden her resim 14 cm'ye zorlanıyordu: 16px'lik bir simge dev, uzun
    bir ekran görüntüsü ise sayfadan taşan bir resim oluyordu."""
    from docx.shared import Cm

    for aday in (veri, None):
        if aday is None:
            aday = _png_e_cevir(veri)  # WEBP/EMF gibi desteklenmeyenler
            if aday is None:
                return False
        olcu = _resim_olcusu(aday)
        try:
            if olcu:
                gen_cm, yuk_cm = olcu[0] / 96 * 2.54, olcu[1] / 96 * 2.54
                olcek = min(1.0, 16 / gen_cm, 22 / yuk_cm)
                belge.add_picture(io.BytesIO(aday), width=Cm(gen_cm * olcek),
                                  height=Cm(yuk_cm * olcek))
            else:
                belge.add_picture(io.BytesIO(aday), width=Cm(14))
            return True
        except Exception:
            continue
    return False


def yaz_pdf(bloklar: list[Blok], yol: Path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image as RLImage
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    # Türkçe destekli yazı tipi ailesi
    font = "Helvetica"
    if Path(_ARIAL).exists():
        kayitli = pdfmetrics.getRegisteredFontNames()
        if "TrArial" not in kayitli:
            pdfmetrics.registerFont(TTFont("TrArial", _ARIAL))
            for ad, dosya in (("TrArial-Bold", _ARIAL_KALIN),
                              ("TrArial-Italic", _ARIAL_ITALIK),
                              ("TrArial-BoldItalic", _ARIAL_KI)):
                if Path(dosya).exists():
                    pdfmetrics.registerFont(TTFont(ad, dosya))
            from reportlab.pdfbase.pdfmetrics import registerFontFamily
            registerFontFamily("TrArial", normal="TrArial",
                               bold="TrArial-Bold", italic="TrArial-Italic",
                               boldItalic="TrArial-BoldItalic")
        font = "TrArial"

    boyutlar = {"h1": 20, "h2": 17, "h3": 15, "h4": 13, "h5": 12, "h6": 11}
    stiller = {
        tip: ParagraphStyle(tip, fontName=font, fontSize=boyut,
                            leading=boyut * 1.35, spaceBefore=10, spaceAfter=6)
        for tip, boyut in boyutlar.items()
    }
    stiller["p"] = ParagraphStyle("p", fontName=font, fontSize=11, leading=15)
    stiller["li"] = ParagraphStyle(
        "li", fontName=font, fontSize=11, leading=15,
        leftIndent=16, bulletIndent=4)

    belge = SimpleDocTemplate(
        str(yol), pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm)
    # Çerçevenin kullanılabilir alanı (6pt iç boşluklar düşülerek)
    azami_gen, azami_yuk = belge.width - 12, belge.height - 24
    akis = []
    for blok in bloklar:
        if blok.tip == "resim":
            # Normalleştir: reportlab'in okuyamadığı biçimler (EMF, WEBP,
            # CMYK/16-bit...) belge oluşturulurken tüm dönüşümü çökertebiliyordu.
            hazir = _pdf_icin_resim(blok.resim)
            if hazir:
                veri, olcu = hazir
                genislik, yukseklik = olcu[0] * 0.75, olcu[1] * 0.75  # px -> pt (96 DPI)
                # Hem genişliğe hem yüksekliğe sığdır: uzun bir resim eskiden
                # "Flowable too large" hatasıyla tüm dönüşümü çökertiyordu.
                olcek = min(1.0, azami_gen / genislik, azami_yuk / yukseklik)
                akis.append(RLImage(io.BytesIO(veri), width=genislik * olcek,
                                    height=yukseklik * olcek))
                akis.append(Spacer(1, 8))
            else:
                akis.append(Paragraph(
                    html.escape(f"[Resim: {guvenli_dosya_adi(blok.resim_ad)}]"),
                    stiller["p"]))
            continue
        ic = ""
        for parca in blok.parcalar:
            m = html.escape(parca.metin)
            if blok.tip.startswith("h") or parca.kalin:
                m = f"<b>{m}</b>"
            if parca.italik:
                m = f"<i>{m}</i>"
            ic += m
        stil = stiller.get(blok.tip, stiller["p"])
        if not ic.strip():
            akis.append(Spacer(1, 8))
        elif blok.tip == "li":
            akis.append(Paragraph(ic, stil, bulletText="•"))
        else:
            akis.append(Paragraph(ic, stil))
    if not akis:
        akis.append(Spacer(1, 8))
    belge.build(akis)


def _rtf_kacir(metin: str) -> str:
    sonuc = []
    for karakter in metin:
        if karakter in "\\{}":
            sonuc.append("\\" + karakter)
        elif karakter == "\t":
            sonuc.append("\\tab ")
        elif karakter == "\n":
            sonuc.append("\\line ")
        elif karakter == "\r":
            continue
        elif ord(karakter) > 127:
            # RTF \uN parametresi işaretli 16-bit'tir: 32767 üstü negatif
            # yazılır, BMP dışı karakterler (emoji) UTF-16 vekil çiftine
            # bölünür. Eskiden "\u128512?" gibi geçersiz kod üretiliyordu.
            kod = karakter.encode("utf-16-le")
            for i in range(0, len(kod), 2):
                birim = int.from_bytes(kod[i:i + 2], "little")
                sonuc.append(f"\\u{birim - 65536 if birim > 32767 else birim}?")
        else:
            sonuc.append(karakter)
    return "".join(sonuc)


def yaz_rtf(bloklar: list[Blok], yol: Path):
    boyutlar = {"h1": 40, "h2": 34, "h3": 30, "h4": 26, "h5": 24, "h6": 22}
    satirlar = [r"{\rtf1\ansi\deff0{\fonttbl{\f0\fcharset162 Arial;}}"]
    for blok in bloklar:
        if blok.tip == "resim":
            satirlar.append(r"\pard\f0\fs22 " + _rtf_kacir(
                f"[Resim: {guvenli_dosya_adi(blok.resim_ad)}]") + r"\par")
            continue
        ic = ""
        for parca in blok.parcalar:
            m = _rtf_kacir(parca.metin)
            if parca.kalin:
                m = r"\b " + m + r"\b0 "
            if parca.italik:
                m = r"\i " + m + r"\i0 "
            ic += m
        if blok.tip.startswith("h"):
            satirlar.append(
                rf"\pard\f0\fs{boyutlar[blok.tip]}\b " + ic + r"\b0\par")
        elif blok.tip == "li":
            satirlar.append(r"\pard\f0\fs22 \bullet  " + ic + r"\par")
        else:
            satirlar.append(r"\pard\f0\fs22 " + ic + r"\par")
    satirlar.append("}")
    yol.write_text("\n".join(satirlar), encoding="ascii")


def yaz_odt(bloklar: list[Blok], yol: Path):
    import mimetypes

    from odf import draw, teletype
    from odf import text as odf_text
    from odf.opendocument import OpenDocumentText
    from odf.style import Style, TextProperties

    belge = OpenDocumentText()
    stil_kalin = Style(name="Kalin", family="text")
    stil_kalin.addElement(TextProperties(fontweight="bold"))
    stil_italik = Style(name="Italik", family="text")
    stil_italik.addElement(TextProperties(fontstyle="italic"))
    stil_ki = Style(name="KalinItalik", family="text")
    stil_ki.addElement(TextProperties(fontweight="bold", fontstyle="italic"))
    for stil in (stil_kalin, stil_italik, stil_ki):
        belge.automaticstyles.addElement(stil)

    def parcalari_ekle(kap, parcalar):
        # teletype: ardışık boşluk, sekme ve satır sonlarını ODF'nin kendi
        # öğelerine (text:s, text:tab, text:line-break) çevirir; düz metin
        # olarak eklenince bunlar kayboluyordu.
        for parca in parcalar:
            if parca.kalin and parca.italik:
                hedef = odf_text.Span(stylename=stil_ki)
            elif parca.kalin:
                hedef = odf_text.Span(stylename=stil_kalin)
            elif parca.italik:
                hedef = odf_text.Span(stylename=stil_italik)
            else:
                teletype.addTextToElement(kap, parca.metin)
                continue
            teletype.addTextToElement(hedef, parca.metin)
            kap.addElement(hedef)

    for blok in bloklar:
        if blok.tip == "resim":
            olcu = _resim_olcusu(blok.resim)
            if olcu:
                gen_cm, yuk_cm = olcu[0] / 37.8, olcu[1] / 37.8
                olcek = min(1.0, 16 / gen_cm, 22 / yuk_cm)
                ad = guvenli_dosya_adi(blok.resim_ad)
                tur = mimetypes.guess_type(ad)[0] or "image/png"
                href = belge.addPictureFromString(blok.resim, tur)
                p = odf_text.P()
                cerceve = draw.Frame(width=f"{gen_cm * olcek:.2f}cm",
                                     height=f"{yuk_cm * olcek:.2f}cm",
                                     anchortype="as-char")
                cerceve.addElement(draw.Image(href=href))
                p.addElement(cerceve)
                belge.text.addElement(p)
            else:
                belge.text.addElement(odf_text.P(
                    text=f"[Resim: {guvenli_dosya_adi(blok.resim_ad)}]"))
            continue
        if blok.tip.startswith("h"):
            eleman = odf_text.H(outlinelevel=int(blok.tip[1]))
        else:
            eleman = odf_text.P()
        parcalari_ekle(eleman, blok.parcalar)
        if blok.tip == "li":
            liste = odf_text.List()
            madde = odf_text.ListItem()
            madde.addElement(eleman)
            liste.addElement(madde)
            belge.text.addElement(liste)
        else:
            belge.text.addElement(eleman)
    belge.save(str(yol))


YAZICILAR = {
    "txt": yaz_txt,
    "md": yaz_md,
    "html": yaz_html,
    "docx": yaz_docx,
    "pdf": yaz_pdf,
    "rtf": yaz_rtf,
    "odt": yaz_odt,
}
